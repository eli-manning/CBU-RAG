"""
CBU Knowledge Base Ingestion
Chunks and embeds CBU content into ChromaDB.

Usage:
    python ingest.py --url https://www.calbaptist.edu/academics/
    python ingest.py --seed          # scrape all default CBU URLs
    python ingest.py --pdf ./docs/handbook.pdf
    python ingest.py --dir ./docs/     # .txt, .pdf, .docx, .xlsx
    python ingest.py --reset         # wipe collection and start fresh
"""

import argparse
import hashlib
import re
import time
from pathlib import Path

import chromadb
import ollama
import requests
from bs4 import BeautifulSoup

# --- Config ---
EMBED_MODEL = "nomic-embed-text"
# nomic-embed-text is trained with task prefixes; omitting them measurably
# degrades retrieval. Documents and queries must use different prefixes.
EMBED_DOC_PREFIX = "search_document: "
CHUNK_SIZE = 220           # words -- smaller chunks retrieve far more precisely
CHUNK_OVERLAP = 40
MIN_CHUNK_WORDS = 20
CHROMA_HOST = "localhost"
CHROMA_PORT = 8001

CBU_SEED_URLS = [
    "https://www.calbaptist.edu/academics/",
    "https://www.calbaptist.edu/admissions/",
    "https://calbaptist.edu/life-at-cbu/",
    "https://www.calbaptist.edu/about/",
    "https://www.calbaptist.edu/engineering/",
    "https://calbaptist.edu/academics/programs/bachelor-of-science-computer-science/",
    "https://calbaptist.edu/academics/programs/minor-computer-science",
    "https://calbaptist.edu/academics/programs/minor-computer-engineering",
    "https://calbaptist.edu/academics/programs/master-of-science-computer-science/",
    "https://calbaptist.edu/academics/programs/minor-data-sciences",
    "https://calbaptist.edu/academics/programs/bachelor-of-science-data-science/",
    "https://calbaptist.edu/academics/programs/bachelor-of-science-electrical-and-computer-engineering/",
    "https://calbaptist.edu/academics/programs/minor-software-engineering-and-app-development"
]

client = chromadb.HttpClient(host=CHROMA_HOST, port=CHROMA_PORT)
collection = client.get_or_create_collection(
    name="cbu_knowledge",
    metadata={"hnsw:space": "cosine"}
)


def _title_for(source: str) -> str:
    """Human-readable document name, used as a per-chunk context header."""
    name = Path(source).name if not source.startswith("http") else source
    for ext in (".txt", ".pdf", ".docx", ".xlsx"):
        name = name.replace(ext, "")
    return name.replace("_", " ").replace("-", " ").strip()


def _split_paragraphs(text: str) -> list[str]:
    """Split on blank lines, falling back to single newlines for flat text."""
    parts = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(parts) <= 1:
        parts = [p.strip() for p in text.split("\n") if p.strip()]
    return parts or [text.strip()]


def chunk_text(text: str, source: str) -> list[dict]:
    """
    Chunk on paragraph boundaries, packing up to CHUNK_SIZE words.

    Splitting mid-sentence on a fixed word window scatters a single fact across
    two chunks and neither retrieves well. Every chunk also carries the document
    title so an isolated passage still says what it belongs to.
    """
    title = _title_for(source)
    paragraphs = _split_paragraphs(text)

    chunks: list[dict] = []
    buf: list[str] = []
    buf_words = 0

    def flush():
        nonlocal buf, buf_words
        if buf_words >= MIN_CHUNK_WORDS:
            body = " ".join(buf).strip()
            chunks.append({
                "text": f"{title}\n\n{body}",
                "source": source,
                "title": title,
                "index": len(chunks),
            })
        buf, buf_words = [], 0

    for para in paragraphs:
        words = para.split()
        if buf_words and buf_words + len(words) > CHUNK_SIZE:
            tail = " ".join(buf).split()[-CHUNK_OVERLAP:]
            flush()
            buf, buf_words = [" ".join(tail)], len(tail)
        buf.append(para)
        buf_words += len(words)
        while buf_words > CHUNK_SIZE:  # a single oversized paragraph
            words_all = " ".join(buf).split()
            head, rest = words_all[:CHUNK_SIZE], words_all[CHUNK_SIZE - CHUNK_OVERLAP:]
            buf, buf_words = [" ".join(head)], len(head)
            flush()
            buf, buf_words = [" ".join(rest)], len(rest)
    flush()
    return chunks


def embed_and_store(chunks: list[dict]):
    for chunk in chunks:
        doc_id = hashlib.md5(chunk["text"].encode()).hexdigest()
        embedding = ollama.embeddings(
            model=EMBED_MODEL, prompt=EMBED_DOC_PREFIX + chunk["text"]
        )["embedding"]
        collection.upsert(
            ids=[doc_id],
            embeddings=[embedding],
            documents=[chunk["text"]],
            metadatas=[{
                "source": chunk["source"],
                "title": chunk.get("title", ""),
                "index": chunk.get("index", 0),
            }],
        )
    print(f"  Stored {len(chunks)} chunks.")


def scrape_url(url: str):
    print(f"Scraping: {url}")
    try:
        resp = requests.get(url, timeout=10, headers={"User-Agent": "CBU-Bot/1.0"})
        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        text = soup.get_text(separator=" ", strip=True)
        chunks = chunk_text(text, source=url)
        embed_and_store(chunks)
        time.sleep(0.5)
    except Exception as e:
        print(f"  Error: {e}")


def ingest_pdf(pdf_path: str):
    try:
        import pymupdf
        print(f"Ingesting PDF: {pdf_path}")
        doc = pymupdf.open(pdf_path)
        text = " ".join(page.get_text() for page in doc)
        chunks = chunk_text(text, source=pdf_path)
        embed_and_store(chunks)
    except ImportError:
        print("Install pymupdf: pip install pymupdf")


def ingest_docx(docx_path: str):
    try:
        import docx
        print(f"Ingesting DOCX: {docx_path}")
        document = docx.Document(docx_path)
        parts = [para.text for para in document.paragraphs if para.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        chunks = chunk_text("\n".join(parts), source=docx_path)
        embed_and_store(chunks)
    except ImportError:
        print("Install python-docx: pip install python-docx")


def ingest_xlsx(xlsx_path: str):
    try:
        import openpyxl
        print(f"Ingesting XLSX: {xlsx_path}")
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        chunks = chunk_text("\n".join(parts), source=xlsx_path)
        embed_and_store(chunks)
    except ImportError:
        print("Install openpyxl: pip install openpyxl")


def ingest_directory(dir_path: str, skip_exts: tuple[str, ...] = ()):
    handlers = {".pdf": ingest_pdf, ".docx": ingest_docx, ".xlsx": ingest_xlsx}
    for path in sorted(Path(dir_path).rglob("*")):
        if not path.is_file() or path.name.startswith("~$") or path.name == ".DS_Store":
            continue
        if path.suffix in skip_exts:
            print(f"  Skipping {path.suffix} (converted separately): {path.name}")
            continue
        if path.suffix == ".txt":
            print(f"Ingesting: {path}")
            text = path.read_text(errors="ignore")
            embed_and_store(chunk_text(text, str(path)))
        elif path.suffix in handlers:
            handlers[path.suffix](str(path))
        else:
            print(f"  Skipping unsupported file: {path.name}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", help="Single URL to scrape")
    parser.add_argument("--seed", action="store_true", help="Scrape all seed URLs")
    parser.add_argument("--pdf", help="PDF to ingest")
    parser.add_argument("--dir", help="Directory of docs to ingest")
    parser.add_argument("--skip-ext", nargs="*", default=[],
                        help="Extensions to skip, e.g. --skip-ext .docx .xlsx")
    parser.add_argument("--reset", action="store_true", help="Wipe collection and exit")
    args = parser.parse_args()

    if args.reset:
        client.delete_collection("cbu_knowledge")
        print("Collection wiped.")
    elif args.url:
        scrape_url(args.url)
    elif args.seed:
        for url in CBU_SEED_URLS:
            scrape_url(url)
    elif args.pdf:
        ingest_pdf(args.pdf)
    elif args.dir:
        ingest_directory(args.dir, tuple(args.skip_ext))
    else:
        parser.print_help()

    print(f"Total docs in collection: {collection.count()}")
